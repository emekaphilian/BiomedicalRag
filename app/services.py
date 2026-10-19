"""Dashboard adapter for the repository's grounded biomedical RAG pipeline."""

from io import BytesIO
from pathlib import Path
import re
from typing import Iterator, List, Optional, Tuple

from rag.document import BiomedicalDocument
from rag.pipeline import BiomedicalRAGPipeline

BASE_CORPUS_DIR = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "processed"
    / "downloaded_sources"
)


def load_default_corpus() -> List[BiomedicalDocument]:
    """Load the repository's prepared biomedical corpus when it is available."""
    if not BASE_CORPUS_DIR.is_dir():
        return []

    from rag.ingestion import load_normalized_documents

    return load_normalized_documents(BASE_CORPUS_DIR)


class RAGService:
    """Expose the upload and query methods expected by the Streamlit UI."""

    def __init__(self, provider: str, model: str, api_key: str):
        provider = provider.lower()
        if provider not in {"openai", "cohere", "huggingface"}:
            raise ValueError(f"Unsupported LLM provider: {provider}")
        if not api_key or not api_key.strip():
            raise ValueError("An API key is required to initialize the LLM.")

        self.provider = provider
        self.model = model
        self._documents = load_default_corpus()
        self.pipeline = BiomedicalRAGPipeline(
            {
                "generator": {
                    "provider": provider,
                    "model_name": model,
                    "api_key": api_key,
                    "use_mock": False,
                },
                "retriever": {"retrieval_k": 20},
            }
        )
        if self._documents:
            self.pipeline.index_documents(self._documents)

    @property
    def total_chunks(self) -> int:
        return len(self.pipeline.chunks)

    @property
    def total_documents(self) -> int:
        return len(self._documents)

    @property
    def source_names(self) -> List[str]:
        return sorted({
            str(document.metadata.get("source", "Unknown"))
            for document in self._documents
        })

    def refresh_default_corpus(self) -> bool:
        """Load the prepared corpus if this cached service has no index yet.

        Streamlit can keep a service in its session after the local corpus has
        been downloaded or generated. Refreshing an empty service here avoids
        making users upload a document just to unlock the prepared corpus.
        """
        if self.pipeline.vectorstore is not None:
            return False

        documents = load_default_corpus()
        if not documents:
            return False

        self._documents = documents
        self.pipeline.index_documents(self._documents)
        return True

    def ingest(self, file_bytes: bytes, filename: str) -> int:
        """Extract and index one uploaded PDF, DOCX, or UTF-8 text file."""
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf":
            from pypdf import PdfReader

            content = "\n".join(
                page.extract_text() or ""
                for page in PdfReader(BytesIO(file_bytes)).pages
            )
        elif suffix == ".docx":
            from docx import Document

            document = Document(BytesIO(file_bytes))
            content = "\n".join(paragraph.text for paragraph in document.paragraphs)
        elif suffix == ".txt":
            content = file_bytes.decode("utf-8-sig")
        else:
            raise ValueError("Supported document formats are PDF, DOCX, and TXT.")

        if not content.strip():
            raise ValueError(f"No readable text was found in {filename}.")

        replacement = BiomedicalDocument(
            content=content,
            metadata={"source": filename, "file": filename},
        )
        existing_index = next(
            (
                index for index, document in enumerate(self._documents)
                if document.metadata.get("source") == filename
            ),
            None,
        )
        if existing_index is None:
            self._documents.append(replacement)
        else:
            self._documents[existing_index] = replacement
        self.pipeline.index_documents(self._documents)
        return len(self.pipeline.chunks)

    def query_stream(
        self,
        question: str,
        chat_history: Optional[list] = None,
        top_k: int = 5,
        mode: str = "hybrid",
    ) -> Tuple[Iterator[str], list]:
        """Run a checked RAG query and return answer chunks plus provenance."""
        chunks, sources, _ = self.query_with_details(
            question,
            chat_history=chat_history,
            top_k=top_k,
            mode=mode,
        )
        return chunks, sources

    def query_with_details(
        self,
        question: str,
        chat_history: Optional[list] = None,
        top_k: int = 5,
        mode: str = "hybrid",
    ) -> Tuple[Iterator[str], list, dict]:
        """Return answer chunks, provenance, and checked pipeline details."""
        del chat_history  # Conversation history is not sent into the medical prompt.
        if mode not in {"hybrid", "dense", "sparse"}:
            raise ValueError(f"Unsupported retrieval mode: {mode}")

        response = self.pipeline.query(
            question,
            top_k=top_k,
            retrieval_mode=mode,
            use_reranking=False,
        )
        sources = [dict(source) for source in response.sources]
        chunks = re.findall(r"\S+\s*", response.answer)
        details = {
            "is_grounded": bool(getattr(response, "is_grounded", False)),
            "retrieval_grounded": bool(
                getattr(response, "retrieval_grounded", getattr(response, "is_grounded", False))
            ),
            "generation_success": bool(
                getattr(response, "generation_success", True)
            ),
            "validation_passed": bool(
                getattr(response, "validation_passed", not getattr(response, "warnings", []))
            ),
            "warnings": list(getattr(response, "warnings", [])),
            "retrieval_time_ms": float(
                getattr(response, "retrieval_time_ms", 0.0)
            ),
            "generation_time_ms": float(
                getattr(response, "generation_time_ms", 0.0)
            ),
        }
        return iter(chunks), sources, details

    def retrieve(
        self,
        question: str,
        top_k: int = 5,
        mode: str = "hybrid",
    ) -> list:
        """Retrieve and return ranked source evidence without generating an answer."""
        if mode not in {"hybrid", "dense", "sparse"}:
            raise ValueError(f"Unsupported retrieval mode: {mode}")
        if top_k < 1:
            raise ValueError("top_k must be at least 1.")
        if not isinstance(question, str) or not question.strip():
            raise ValueError("The query must be a non-empty string.")
        if self.pipeline.vectorstore is None:
            raise RuntimeError("Index documents before running retrieval.")

        from rag.safety import process_safe_query

        safety_result = process_safe_query(question)
        if not safety_result.allowed:
            raise ValueError(safety_result.response)

        if mode == "hybrid":
            from rag.hybrid_retriever import HybridRetriever

            response = HybridRetriever(
                self.pipeline.vectorstore,
                self.pipeline.chunks,
            ).retrieve(safety_result.query, k=top_k)
            retrieved = response["fused_results"][:top_k]
            metadata_by_content = {}
            for result in response["dense_results"]:
                metadata_by_content.setdefault(result.content, result.metadata)
            for result in response["sparse_results"]:
                document = result["document"]
                metadata_by_content.setdefault(document.content, document.metadata)
            candidates = [
                (content, metadata_by_content[content], score)
                for content, score in retrieved
                if content in metadata_by_content
            ]
        elif mode == "dense":
            from rag.retriever import BiomedicalRetriever

            response = BiomedicalRetriever(self.pipeline.vectorstore).retrieve(
                safety_result.query,
                k=top_k,
            )
            candidates = [
                (
                    result.content,
                    result.metadata,
                    result.similarity_score,
                )
                for result in response["results"]
            ]
        else:
            from rag.bm25_retriever import BM25Retriever

            response = BM25Retriever(self.pipeline.documents).search(
                safety_result.query,
                k=top_k,
            )
            candidates = [
                (
                    result["document"].content,
                    result["document"].metadata,
                    result["score"],
                )
                for result in response
            ]

        return [
            {
                "content": content,
                "metadata": dict(metadata),
                "score": float(score),
                "rank": rank,
            }
            for rank, (content, metadata, score) in enumerate(candidates, 1)
        ]
